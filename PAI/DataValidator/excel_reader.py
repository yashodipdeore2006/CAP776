from openpyxl import load_workbook


# reading excel file
def read_data(file_path):

    # loading excel workbook
    workbook = load_workbook(file_path, data_only=True)

    # selecting daily log sheet
    sheet = workbook["Daily Log"]

    # storing excel sheet data
    data = []

    # reading each row from excel sheet
    for row in sheet.iter_rows(min_row=7, values_only=True):

        # stopping when there is no date
        if row[0] is None:
            break

        # storing each day data in the form of object
        day_data = {
            "date": row[0],
            "sleep": row[1],
            "fitness": row[2],
            "study": row[3],
            "coding": row[4],
            "class": row[5],
            "classes_attended": row[6],
            "other_activities": row[7],
            "total_tracked": row[8],
            "free_time": row[9],
            "feeling": row[10],
            "satisfaction": row[11],
            "energy": row[12],
            "notes": row[13]
        }

        # adding day data to the main data
        data.append(day_data)

    return data